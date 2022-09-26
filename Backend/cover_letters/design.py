from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, RGBColor, Pt
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from para import add_para


def template_3():
    # Declaring the document variable
    document = Document()

    # Defining Sections of the document
    sections = document.sections

    # Setting margins
    for section in sections:
        section.top_margin = Inches(0)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # Table
    table = document.add_table(2, 2)
    table.autofit = False
    table.cell(1, 0).merge(table.cell(1, 1))
    table.cell(0, 1).width = Inches(3)
    
    #To remove the additional line in the cells
    for cell in table._cells:
        cell._element.clear_content()
    
    
    data = ['<Date>', '<Company Address>', 'Dear <Addressing-To>', 'In my pursuit of new creative endeavors, was excited to find the Costume Designer opportunity with the International Peace & Film Festival in collaboration with NBC New York. As a progressive professional with styling expertise and fashion design knowledge, believe that I can bring valuable contributions to the festivals fashion team and exhibition', 'In the past, I have designed dynamic pieces for film and television as well as theater it has been a huge development in my career to work with diverse elements for each scenario and be able to showcase scenes and settings through the presence of fashion. take my inspiration from multiple decades and sources across the world, these varying concepts have given my garments the recognition of being unique, daring, and captivating My philosophy is that introducing fresh perspectives and new techniques allow for an open and creative environment. Similar to the International Peace & Film Festival, my goal is to attract and illuminate viewers and festival-goors on the beauty of fashion and textiles', 'In my Fashion Designer role, my successes have stemmed from essential skills, including pattem development and multi print assembles. am a collaborative team player who leads a team of eight designers and is always searching for opportunities to impart valuable insights into them as a team and a collaborator. With these quales, I am able to facilitate positive impact in the fashion world', 'For a greater presentation of my background and views of my designs, textile compilations, and qualifications please review my attached resume and link to my website. am eager to further speak with you about the exhibition and greatly appreciate your consideration', 'Thank you for your time', "name.full_name"]
    
    #Adding Name
    add_para(
        "name['full_name']",
        cell=table.cell(1, 0),
        document=document,
        font_size=60,
        line_spacing=0.6,
        font_name="Calibri Light",
        color=RGBColor(26, 23, 24),
        style_name="full_name",
        font_bold=True,
    )
    add_para(
        "contact['number']",
        cell=table.cell(0, 1),
        document=document,
        font_size=20,
        line_spacing=0.6,
        font_name="Calibri Light",
        color=RGBColor(26, 23, 24),
        style_name="number",
        font_bold=True,
    )
    add_para(
        "contact['address']",
        cell=table.cell(0, 1),
        document=document,
        font_size=20,
        line_spacing=0.6,
        font_name="Calibri Light",
        color=RGBColor(26, 23, 24),
        style_name="address",
        font_bold=True,
    )
    add_para(
        "contact['email']",
        cell=table.cell(0, 1),
        document=document,
        font_size=20,
        line_spacing=0.6,
        font_name="Calibri Light",
        color=RGBColor(26, 23, 24),
        style_name="email",
        font_bold=True,
    )
    
    #To increase the height of the mentioned cell 
    add_para(
        " ",
        cell=table.cell(0, 0),
        document=document,
        style_name="space",
    )
    i = 0
    for parag in data:
        add_para(
            parag,
            cell=table.cell(1, 0),
            document=document,
            font_size=20,
            line_spacing=0.6,
            font_name="Calibri Light",
            color=RGBColor(26, 23, 24),
            style_name=str(i),
            font_bold=True,
        )
        i+=1
    
    # Coloring Cells
    shading_elm_1 = parse_xml(r'<w:shd {} w:fill="262324"/>'.format(nsdecls("w")))
    table.rows[0].cells[0]._tc.get_or_add_tcPr().append(shading_elm_1)
    shading_elm_2 = parse_xml(r'<w:shd {} w:fill="262324"/>'.format(nsdecls("w")))
    table.rows[2].cells[0]._tc.get_or_add_tcPr().append(shading_elm_2)

    # Saving the document
    document.save("design.docx")


template_3()
